import csv, io, json
from datetime import datetime
from decimal import Decimal, InvalidOperation
from fastapi import HTTPException
from openpyxl import load_workbook

REQUIRED = ["date", "product_name", "quantity", "unit_price", "unit_cost"]
LIMIT = 5 * 1024 * 1024


def parse_file(content, filename, mapping=None):
    if len(content) > LIMIT:
        raise HTTPException(413, detail="file_too_large")
    try:
        if filename.lower().endswith(".csv"):
            text = content.decode("utf-8-sig")
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            rows = list(csv.reader(io.StringIO(text), dialect))
        elif filename.lower().endswith(".xlsx"):
            import zipfile

            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                    raise ValueError()
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = []
            for row in wb.active.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > 10001:
                    raise HTTPException(422, detail="too_many_rows")
            wb.close()
        else:
            raise HTTPException(415, detail="unsupported_format")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, detail="invalid_file")
    if len(rows) < 2:
        raise HTTPException(422, detail="empty_file")
    if len(rows) > 10001:
        raise HTTPException(422, detail="too_many_rows")
    headers = [str(h).strip() for h in rows[0]]
    if len(set(headers)) != len(headers):
        raise HTTPException(422, detail="duplicate_columns")
    mapping = mapping or {k: k for k in REQUIRED}
    missing = [k for k in REQUIRED if mapping.get(k) not in headers]
    if missing:
        return {
            "headers": headers,
            "missing": missing,
            "valid_rows": [],
            "errors": [],
            "preview": [],
        }
    indexes = {k: headers.index(mapping[k]) for k in REQUIRED}
    valid = []
    errors = []
    for n, row in enumerate(rows[1:], 2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        try:
            raw = {k: row[i] for k, i in indexes.items()}
            when = (
                raw["date"]
                if isinstance(raw["date"], datetime)
                else datetime.fromisoformat(str(raw["date"]))
            )
            if when.tzinfo:
                when = when.replace(tzinfo=None)
            q = Decimal(str(raw["quantity"]))
            price = Decimal(str(raw["unit_price"]))
            cost = Decimal(str(raw["unit_cost"]))
            name = str(raw["product_name"] or "").strip()
            if (
                not name
                or len(name) > 160
                or not q.is_finite()
                or q != q.to_integral()
                or not 0 < q <= 1000000
                or not price.is_finite()
                or not cost.is_finite()
                or not 0 < price <= 10000000
                or not 0 <= cost <= 10000000
            ):
                raise ValueError()
            if price.as_tuple().exponent < -2 or cost.as_tuple().exponent < -2:
                raise ValueError()
            valid.append(
                {
                    "date": when.isoformat(),
                    "product_name": name,
                    "quantity": int(q),
                    "unit_price": str(price),
                    "unit_cost": str(cost),
                }
            )
        except (ValueError, InvalidOperation, IndexError, TypeError):
            errors.append({"row": n, "code": "invalid_row"})
    return {
        "headers": headers,
        "missing": [],
        "valid_rows": valid,
        "preview": valid[:8],
        "errors": errors,
        "valid_count": len(valid),
        "rejected_count": len(errors),
    }
